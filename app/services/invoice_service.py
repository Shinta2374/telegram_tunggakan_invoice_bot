from pathlib import Path
from threading import Lock

from openpyxl import load_workbook


class InvoiceService:

    _cache_lock = Lock()
    _cache_mtime = None
    _invoice_by_customer = {}
    _invoice_customers = []
    _invoice_period = ""

    def __init__(self):
        self.file_path = (
            Path("data")
            / "invoice"
            / "invoice.xlsx"
        )

    def file_exists(self):
        return self.file_path.exists()

    def _normalize_text(self, value):
        if value is None:
            return ""

        return (
            str(value)
            .replace("\xa0", " ")
            .strip()
        )

    def _normalize_customer_id(self, value):
        if value is None:
            return ""

        text = (
            str(value)
            .replace("\xa0", "")
            .strip()
        )

        if not text:
            return ""

        if text.endswith(".0"):
            try:
                number = float(text)

                if number.is_integer():
                    return str(int(number))

            except (
                ValueError,
                TypeError,
            ):
                pass

        try:
            number = float(text)

            if number.is_integer():
                return str(int(number))

        except (
            ValueError,
            TypeError,
        ):
            pass

        return text

    def _to_number(self, value):
        if value is None:
            return 0

        if isinstance(value, bool):
            return 0

        if isinstance(value, (int, float)):
            try:
                if value != value:
                    return 0

                return float(value)

            except Exception:
                return 0

        text = self._normalize_text(value)

        if not text:
            return 0

        text = (
            text
            .replace("Rp", "")
            .replace("rp", "")
            .replace(" ", "")
        )

        if "." in text and "," not in text:
            parts = text.split(".")

            if all(
                part.isdigit()
                for part in parts
            ):
                if all(
                    len(part) == 3
                    for part in parts[1:]
                ):
                    text = "".join(parts)

        elif "." in text and "," in text:
            text = (
                text
                .replace(".", "")
                .replace(",", ".")
            )

        elif "," in text:
            text = text.replace(",", ".")

        try:
            return float(text)

        except (
            ValueError,
            TypeError,
        ):
            return 0

    def _normalize_status(self, value):
        status = self._normalize_text(
            value
        )

        if not status:
            return "On Progress"

        status_lower = status.lower()

        if (
            "no data to display"
            in status_lower
        ):
            return "On Progress"

        if (
            status_lower == "#n/a"
            or status_lower == "n/a"
            or "#n/a" in status_lower
        ):
            return "On Progress"

        manual_keywords = [
            "inv manual",
            "invoice manual",
            "sent tghn manual",
            "sent manual",
            "manual via ideas",
            "klik sent manual",
        ]

        for keyword in manual_keywords:
            if keyword in status_lower:
                return "Invoice Manual"

        if (
            "message has been sent"
            in status_lower
            or status_lower == "sent"
            or "terkirim" in status_lower
        ):
            return "Terkirim"

        return status

    def _format_period(self, value):
        text = self._normalize_text(value)

        if not text:
            return ""

        if text.endswith(".0"):
            try:
                number = float(text)

                if number.is_integer():
                    text = str(int(number))

            except (
                ValueError,
                TypeError,
            ):
                pass

        if (
            len(text) == 6
            and text.isdigit()
        ):
            year = text[:4]
            month = text[4:6]

            month_names = {
                "01": "Januari",
                "02": "Februari",
                "03": "Maret",
                "04": "April",
                "05": "Mei",
                "06": "Juni",
                "07": "Juli",
                "08": "Agustus",
                "09": "September",
                "10": "Oktober",
                "11": "November",
                "12": "Desember",
            }

            month_name = month_names.get(
                month
            )

            if month_name:
                return (
                    f"{month_name} "
                    f"{year}"
                )

        return text

    def _find_header_row(self, worksheet):
        required_headers = {
            "no.jastel": None,
            "contr.account detail": None,
            "bill.pe": None,
            "ppn": None,
            "total amount": None,
            "stts": None,
        }

        max_check = min(
            worksheet.max_row,
            20,
        )

        for row_number in range(
            1,
            max_check + 1,
        ):
            found = {}

            for cell in worksheet[row_number]:
                value = self._normalize_text(
                    cell.value
                )

                normalized = (
                    value
                    .lower()
                    .replace(" ", "")
                )

                if normalized == "nojastel":
                    found[
                        "no.jastel"
                    ] = cell.column

                elif (
                    normalized
                    == "contr.accountdetail"
                ):
                    found[
                        "contr.account detail"
                    ] = cell.column

                elif normalized == "bill.pe":
                    found[
                        "bill.pe"
                    ] = cell.column

                elif normalized == "ppn":
                    found[
                        "ppn"
                    ] = cell.column

                elif (
                    normalized
                    == "totalamount"
                ):
                    found[
                        "total amount"
                    ] = cell.column

                elif normalized == "stts":
                    found[
                        "stts"
                    ] = cell.column

            if all(
                key in found
                for key in required_headers
            ):
                found["billing amount"] = (
                    found["ppn"] - 1
                )

                return (
                    row_number,
                    found,
                )

        return None

    def _get_verified_columns(self):
        return {
            "no": 1,
            "no.jastel": 2,
            "contr.account detail": 3,
            "bill.pe": 4,
            "billing amount": 5,
            "ppn": 6,
            "total amount": 7,
            "stts": 8,
        }

    def _get_file_mtime(self):
        if not self.file_exists():
            return None

        try:
            return self.file_path.stat().st_mtime_ns

        except OSError:
            return None

    def _load_workbook(self):
        if not self.file_exists():
            raise FileNotFoundError(
                "File invoice.xlsx "
                "tidak ditemukan."
            )

        try:
            return load_workbook(
                filename=self.file_path,
                data_only=True,
                read_only=True,
            )

        except Exception as error:
            raise ValueError(
                "Gagal membaca "
                "invoice.xlsx: "
                f"{error}"
            )

    def _build_cache(self):
        current_mtime = (
            self._get_file_mtime()
        )

        if current_mtime is None:
            return

        print(
            "[INVOICE CACHE] "
            "Membaca invoice.xlsx..."
        )

        workbook = self._load_workbook()

        invoice_by_customer = {}
        invoice_customers = []
        customer_map = {}
        periods = []

        try:
            for worksheet in workbook.worksheets:

                header_result = (
                    self._find_header_row(
                        worksheet
                    )
                )

                if header_result:
                    (
                        header_row,
                        columns,
                    ) = header_result

                    start_row = (
                        header_row + 1
                    )

                else:
                    columns = (
                        self._get_verified_columns()
                    )

                    start_row = 2

                no_jastel_col = columns[
                    "no.jastel"
                ]

                customer_name_col = columns[
                    "contr.account detail"
                ]

                periode_col = columns[
                    "bill.pe"
                ]

                billing_amount_col = columns[
                    "billing amount"
                ]

                ppn_col = columns[
                    "ppn"
                ]

                total_amount_col = columns[
                    "total amount"
                ]

                status_col = columns[
                    "stts"
                ]

                for row_number in range(
                    start_row,
                    worksheet.max_row + 1,
                ):

                    raw_customer_id = (
                        worksheet.cell(
                            row=row_number,
                            column=no_jastel_col,
                        ).value
                    )

                    customer_id = (
                        self._normalize_customer_id(
                            raw_customer_id
                        )
                    )

                    if not customer_id:
                        continue

                    raw_customer_name = (
                        worksheet.cell(
                            row=row_number,
                            column=customer_name_col,
                        ).value
                    )

                    customer_name = (
                        self._normalize_text(
                            raw_customer_name
                        )
                    )

                    raw_periode = (
                        worksheet.cell(
                            row=row_number,
                            column=periode_col,
                        ).value
                    )

                    periode = (
                        self._normalize_text(
                            raw_periode
                        )
                    )

                    if periode:
                        periods.append(
                            periode
                        )

                    raw_billing = (
                        worksheet.cell(
                            row=row_number,
                            column=billing_amount_col,
                        ).value
                    )

                    raw_ppn = (
                        worksheet.cell(
                            row=row_number,
                            column=ppn_col,
                        ).value
                    )

                    raw_total = (
                        worksheet.cell(
                            row=row_number,
                            column=total_amount_col,
                        ).value
                    )

                    raw_status = (
                        worksheet.cell(
                            row=row_number,
                            column=status_col,
                        ).value
                    )

                    invoice = {
                        "customer_id": customer_id,
                        "pelanggan": customer_name,
                        "customer_name": customer_name,
                        "periode": periode,
                        "billing_amount": (
                            self._to_number(
                                raw_billing
                            )
                        ),
                        "ppn": (
                            self._to_number(
                                raw_ppn
                            )
                        ),
                        "total_amount": (
                            self._to_number(
                                raw_total
                            )
                        ),
                        "status": (
                            self._normalize_status(
                                raw_status
                            )
                        ),
                        "status_raw": (
                            self._normalize_text(
                                raw_status
                            )
                        ),
                        "sheet": worksheet.title,
                        "row": row_number,
                    }

                    if customer_id not in (
                        invoice_by_customer
                    ):
                        invoice_by_customer[
                            customer_id
                        ] = []

                    invoice_by_customer[
                        customer_id
                    ].append(invoice)

                    if customer_id not in customer_map:
                        customer_map[
                            customer_id
                        ] = {
                            "customer_id": customer_id,
                            "customer_name": (
                                customer_name
                                or "-"
                            ),
                            "pelanggan": (
                                customer_name
                                or "-"
                            ),
                        }

            invoice_customers = list(
                customer_map.values()
            )

            invoice_customers.sort(
                key=lambda item: (
                    item.get(
                        "customer_name",
                        "",
                    ).lower()
                )
            )

            invoice_period = ""

            if periods:
                period_counter = {}

                for period in periods:
                    normalized_period = (
                        self._normalize_text(
                            period
                        )
                    )

                    if not normalized_period:
                        continue

                    period_counter[
                        normalized_period
                    ] = (
                        period_counter.get(
                            normalized_period,
                            0,
                        )
                        + 1
                    )

                if period_counter:
                    invoice_period = max(
                        period_counter,
                        key=period_counter.get,
                    )

            InvoiceService._invoice_by_customer = (
                invoice_by_customer
            )

            InvoiceService._invoice_customers = (
                invoice_customers
            )

            InvoiceService._invoice_period = (
                invoice_period
            )

            InvoiceService._cache_mtime = (
                current_mtime
            )

            print(
                "[INVOICE CACHE] "
                f"Cache berhasil dibuat. "
                f"{len(invoice_customers)} "
                "customer, "
                f"{sum(len(items) for items in invoice_by_customer.values())} "
                "invoice."
            )

        finally:
            workbook.close()

    def _ensure_cache(self):
        current_mtime = (
            self._get_file_mtime()
        )

        if current_mtime is None:
            return

        if (
            InvoiceService._cache_mtime
            == current_mtime
            and InvoiceService._invoice_customers
            is not None
        ):
            return

        with InvoiceService._cache_lock:

            current_mtime = (
                self._get_file_mtime()
            )

            if current_mtime is None:
                return

            if (
                InvoiceService._cache_mtime
                == current_mtime
            ):
                return

            self._build_cache()

    def get_invoice_customers(self):
        self._ensure_cache()

        return [
            dict(customer)
            for customer in (
                InvoiceService._invoice_customers
            )
        ]

    def get_customer_invoice_data(
        self,
        customer_id,
    ):
        self._ensure_cache()

        target_id = (
            self._normalize_customer_id(
                customer_id
            )
        )

        if not target_id:
            return []

        invoices = (
            InvoiceService
            ._invoice_by_customer
            .get(
                target_id,
                [],
            )
        )

        return [
            dict(invoice)
            for invoice in invoices
        ]

    def get_invoice_customer(
        self,
        customer_id,
    ):
        self._ensure_cache()

        target_id = (
            self._normalize_customer_id(
                customer_id
            )
        )

        if not target_id:
            return None

        for customer in (
            InvoiceService
            ._invoice_customers
        ):
            if (
                customer.get(
                    "customer_id"
                )
                == target_id
            ):
                return dict(customer)

        return None

    def get_invoice_period(self):
        self._ensure_cache()

        if not (
            InvoiceService
            ._invoice_period
        ):
            return ""

        return self._format_period(
            InvoiceService._invoice_period
        )

    def clear_cache(self):
        with InvoiceService._cache_lock:
            InvoiceService._cache_mtime = None
            InvoiceService._invoice_by_customer = {}
            InvoiceService._invoice_customers = []
            InvoiceService._invoice_period = ""

            print(
                "[INVOICE CACHE] "
                "Cache dibersihkan."
            )